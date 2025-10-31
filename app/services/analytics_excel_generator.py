"""Excel generator for Global Analytics Statistics Reports.

This module generates Excel files for the Statistics tab functionality,
creating reports with data from ALL projects (not individual projects).

The Excel files have 2 sheets:
- Sheet 1: "Project-Specific" - One row per project
- Sheet 2: "Region-Specific" - One row per region (aggregated)

IMPORTANT: This is DIFFERENT from analytics_report_generator.py (which handles
individual project reports using templates).
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.analytics_models import (
    BSRProjectAnalytics,
    BSRRegionAnalytics,
    NWProjectAnalytics,
    NWRegionAnalytics,
)
from app.utils.logging_utils import get_logger

logger = get_logger(__name__)


class AnalyticsExcelGenerator:
    """Generator for analytics Excel reports with proper formatting."""

    # Header styling
    HEADER_FONT = Font(bold=True, size=12, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data cell styling
    DATA_ALIGNMENT_LEFT = Alignment(horizontal="left", vertical="center")
    DATA_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center")
    DATA_ALIGNMENT_RIGHT = Alignment(horizontal="right", vertical="center")

    def _apply_header_style(self, worksheet, row: int, num_columns: int):
        """Apply header styling to the first row of a worksheet.

        Args:
            worksheet: Openpyxl worksheet object
            row: Row number to apply styling to (typically 1)
            num_columns: Number of columns to style
        """
        for col in range(1, num_columns + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT

    def _auto_size_columns(self, worksheet, num_columns: int):
        """Auto-size columns based on content width.

        Args:
            worksheet: Openpyxl worksheet object
            num_columns: Number of columns to resize
        """
        for col in range(1, num_columns + 1):
            max_length = 0
            column_letter = get_column_letter(col)

            for cell in worksheet[column_letter]:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length

            # Set column width (with some padding)
            adjusted_width = min(max_length + 2, 50)  # Max 50 chars
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def generate_nw_analytics_excel(
        self,
        projects: List[NWProjectAnalytics],
        regions: List[NWRegionAnalytics]
    ) -> BytesIO:
        """Generate Excel file for NW Analytics Statistics.

        Creates a workbook with 2 sheets:
        - Project-Specific: All NW projects data
        - Region-Specific: Aggregated region data

        Args:
            projects: List of NW project analytics
            regions: List of NW region analytics

        Returns:
            BytesIO object containing the Excel file
        """
        logger.info(
            "Generating NW Analytics Excel with %d projects and %d regions",
            len(projects),
            len(regions)
        )

        workbook = Workbook()

        # Remove default sheet
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]

        # Sheet 1: Project-Specific
        ws_projects = workbook.create_sheet("Project-Specific", 0)

        # Headers for Project-Specific sheet (matches Excel template exactly)
        project_headers = [
            "Region",
            "Active Lead",
            "Project Name",
            "DisplayName",
            "ClientName",
            "Percent Retained (Positive + Neutral)",
            "Percent Positive",
            "Percent Neutral",
            "Percent Negative",
            "Names Created"
        ]

        # Write headers
        for col, header in enumerate(project_headers, start=1):
            ws_projects.cell(row=1, column=col, value=header)

        # Apply header styling
        self._apply_header_style(ws_projects, 1, len(project_headers))

        # Write project data
        for row, project in enumerate(projects, start=2):
            ws_projects.cell(row=row, column=1, value=project.region).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=2, value=project.active_lead).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=3, value=project.project_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=4, value=project.display_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=5, value=project.client_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=6, value=project.percent_retained).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=7, value=project.percent_positive).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=8, value=project.percent_neutral).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=9, value=project.percent_negative).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=10, value=project.names_created).alignment = self.DATA_ALIGNMENT_RIGHT

        # Auto-size columns
        self._auto_size_columns(ws_projects, len(project_headers))

        # Sheet 2: Region-Specific
        ws_regions = workbook.create_sheet("Region-Specific", 1)

        # Headers for Region-Specific sheet
        region_headers = [
            "Region",
            "Active Lead",
            "Projects to Date",
            "% Retained",
            "% Positive",
            "% Neutral",
            "% Negative",
            "Avg Names per Project"
        ]

        # Write headers
        for col, header in enumerate(region_headers, start=1):
            ws_regions.cell(row=1, column=col, value=header)

        # Apply header styling
        self._apply_header_style(ws_regions, 1, len(region_headers))

        # Write region data
        for row, region in enumerate(regions, start=2):
            ws_regions.cell(row=row, column=1, value=region.region).alignment = self.DATA_ALIGNMENT_LEFT
            ws_regions.cell(row=row, column=2, value=region.active_lead).alignment = self.DATA_ALIGNMENT_LEFT
            ws_regions.cell(row=row, column=3, value=region.projects_to_date).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=4, value=region.percent_retained).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=5, value=region.percent_positive).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=6, value=region.percent_neutral).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=7, value=region.percent_negative).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=8, value=f"{region.avg_names_per_project:.2f}").alignment = self.DATA_ALIGNMENT_RIGHT

        # Auto-size columns
        self._auto_size_columns(ws_regions, len(region_headers))

        # Save to BytesIO
        excel_buffer = BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)

        logger.info("NW Analytics Excel file generated successfully")
        return excel_buffer

    def generate_bsr_analytics_excel(
        self,
        projects: List[BSRProjectAnalytics],
        regions: List[BSRRegionAnalytics]
    ) -> BytesIO:
        """Generate Excel file for BSR Analytics Statistics.

        Creates a workbook with 2 sheets:
        - Project-Specific: All BSR projects data
        - Region-Specific: Aggregated region data

        Args:
            projects: List of BSR project analytics
            regions: List of BSR region analytics

        Returns:
            BytesIO object containing the Excel file
        """
        logger.info(
            "Generating BSR Analytics Excel with %d projects and %d regions",
            len(projects),
            len(regions)
        )

        workbook = Workbook()

        # Remove default sheet
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]

        # Sheet 1: Project-Specific
        ws_projects = workbook.create_sheet("Project-Specific", 0)

        # Headers for Project-Specific sheet (matches Excel template exactly)
        project_headers = [
            "Region",
            "Active Lead",
            "Project Name",
            "DisplayName",
            "ClientName",
            "Names Created (Web [Post-Its + Name Entry Box])",
            "Names Created (Mobile)",
            "Total Names Created"
        ]

        # Write headers
        for col, header in enumerate(project_headers, start=1):
            ws_projects.cell(row=1, column=col, value=header)

        # Apply header styling
        self._apply_header_style(ws_projects, 1, len(project_headers))

        # Write project data
        for row, project in enumerate(projects, start=2):
            ws_projects.cell(row=row, column=1, value=project.region).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=2, value=project.active_lead).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=3, value=project.project_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=4, value=project.display_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=5, value=project.client_name).alignment = self.DATA_ALIGNMENT_LEFT
            ws_projects.cell(row=row, column=6, value=project.names_created_web).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=7, value=project.names_created_mobile).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_projects.cell(row=row, column=8, value=project.total_names_created).alignment = self.DATA_ALIGNMENT_RIGHT

        # Auto-size columns
        self._auto_size_columns(ws_projects, len(project_headers))

        # Sheet 2: Region-Specific
        ws_regions = workbook.create_sheet("Region-Specific", 1)

        # Headers for Region-Specific sheet
        region_headers = [
            "Region",
            "Active Lead",
            "Projects to Date",
            "PC Count",
            "Mobile Count",
            "Total"
        ]

        # Write headers
        for col, header in enumerate(region_headers, start=1):
            ws_regions.cell(row=1, column=col, value=header)

        # Apply header styling
        self._apply_header_style(ws_regions, 1, len(region_headers))

        # Write region data
        for row, region in enumerate(regions, start=2):
            ws_regions.cell(row=row, column=1, value=region.region).alignment = self.DATA_ALIGNMENT_LEFT
            ws_regions.cell(row=row, column=2, value=region.active_lead).alignment = self.DATA_ALIGNMENT_LEFT
            ws_regions.cell(row=row, column=3, value=region.projects_to_date).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=4, value=region.pc_count).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=5, value=region.mobile_count).alignment = self.DATA_ALIGNMENT_RIGHT
            ws_regions.cell(row=row, column=6, value=region.total).alignment = self.DATA_ALIGNMENT_RIGHT

        # Auto-size columns
        self._auto_size_columns(ws_regions, len(region_headers))

        # Save to BytesIO
        excel_buffer = BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)

        logger.info("BSR Analytics Excel file generated successfully")
        return excel_buffer

    def generate_analytics_excel(
        self,
        project_type: str,
        projects_data: List[dict],
        regions_data: List[dict]
    ) -> tuple[BytesIO, str]:
        """Generate analytics Excel file based on project type.

        Args:
            project_type: "NW" or "BSR"
            projects_data: List of project analytics dictionaries
            regions_data: List of region analytics dictionaries

        Returns:
            Tuple of (BytesIO Excel file, filename)

        Raises:
            ValueError: If project_type is invalid
        """
        timestamp = datetime.now().strftime("%Y%m%d")
        filename = f"{project_type}AnalyticsReport_{timestamp}.xlsx"

        if project_type == "NW":
            # Convert dicts to Pydantic models
            projects = [NWProjectAnalytics(**p) for p in projects_data]
            regions = [NWRegionAnalytics(**r) for r in regions_data]
            excel_buffer = self.generate_nw_analytics_excel(projects, regions)
        elif project_type == "BSR":
            # Convert dicts to Pydantic models
            projects = [BSRProjectAnalytics(**p) for p in projects_data]
            regions = [BSRRegionAnalytics(**r) for r in regions_data]
            excel_buffer = self.generate_bsr_analytics_excel(projects, regions)
        else:
            raise ValueError(f"Invalid project_type: {project_type}. Must be 'NW' or 'BSR'")

        return excel_buffer, filename


# Global generator instance
analytics_excel_generator = AnalyticsExcelGenerator()

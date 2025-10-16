"""Service for generating NW Analytics Excel reports."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from app.config.settings import settings
from app.services.nw_reports_service import nw_reports_service
from app.utils.logging_utils import get_logger
from app.utils.nw_data_utils import sanitize_filename
from app.utils.path_utils import get_nw_analytics_dir

logger = get_logger(__name__)

# Template path for NW Analytics - relative to project root
ANALYTICS_TEMPLATE_PATH = settings.app_dir / "templates" / "NW_Analytics" / "nwanalytics.xlsx"


class AnalyticsReportGenerator:
    """Generator for NW Analytics Excel reports."""

    def generate_analytics_report(
        self,
        presentation_id: int,
        project_name: str,
        display_name: str,
    ) -> Path:
        """Generate Analytics Excel report from template.

        Args:
            presentation_id: Presentation ID to generate report for
            project_name: Project name
            display_name: Display name for the presentation

        Returns:
            Path to generated Excel file

        Raises:
            FileNotFoundError: If template file is not found
        """
        logger.info(
            "Starting Analytics report generation for presentation_id=%d",
            presentation_id,
        )

        # Check if template exists
        if not ANALYTICS_TEMPLATE_PATH.exists():
            error_msg = f"Analytics template not found at: {ANALYTICS_TEMPLATE_PATH}"
            logger.error(error_msg)
            raise FileNotFoundError(error_msg)

        # Load template
        workbook = load_workbook(str(ANALYTICS_TEMPLATE_PATH))

        # Fill Project-Specific sheet
        self._fill_project_specific_sheet(workbook, presentation_id)

        # Fill Region-Specific sheet
        self._fill_region_specific_sheet(workbook, presentation_id)

        # Save workbook to NW Analytics directory
        # Format: NWAnalyticsReport_[FechaActual].xlsx
        output_dir = get_nw_analytics_dir()
        output_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d")
        filename = f"NWAnalyticsReport_{timestamp}.xlsx"
        output_path = output_dir / filename

        workbook.save(str(output_path))
        logger.info("Analytics report saved to: %s", output_path)

        return output_path

    def _fill_project_specific_sheet(self, workbook, presentation_id: int) -> None:
        """Fill 'Project-Specific' sheet with analytics data.

        Args:
            workbook: Openpyxl workbook object
            presentation_id: Presentation ID
        """
        logger.debug("Filling Project-Specific sheet")

        # Get or create sheet
        sheet_name = "Project-Specific"
        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
        else:
            ws = workbook.create_sheet(sheet_name)
            logger.warning("Sheet '%s' not found in template, created new sheet", sheet_name)

        # Get analytics data
        analytics = nw_reports_service.get_project_analytics(presentation_id)

        # Clear existing data (keep headers)
        # Assuming data starts from row 2
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        # Write data starting from row 2
        row_num = 2
        for item in analytics:
            ws.cell(row=row_num, column=1, value=item.metric_name)
            ws.cell(row=row_num, column=2, value=item.metric_value)
            row_num += 1

        logger.info("Project-Specific sheet filled with %d metrics", len(analytics))

    def _fill_region_specific_sheet(self, workbook, presentation_id: int) -> None:
        """Fill 'Region-Specific' sheet with analytics data.

        Args:
            workbook: Openpyxl workbook object
            presentation_id: Presentation ID
        """
        logger.debug("Filling Region-Specific sheet")

        # Get or create sheet
        sheet_name = "Region-Specific"
        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
        else:
            ws = workbook.create_sheet(sheet_name)
            logger.warning("Sheet '%s' not found in template, created new sheet", sheet_name)

        # Get analytics data
        analytics = nw_reports_service.get_region_specific_analytics(presentation_id)

        # Clear existing data (keep headers)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        # Write data starting from row 2
        row_num = 2
        for item in analytics:
            ws.cell(row=row_num, column=1, value=item.region)
            ws.cell(row=row_num, column=2, value=item.metric_name)
            ws.cell(row=row_num, column=3, value=item.metric_value)
            row_num += 1

        logger.info("Region-Specific sheet filled with %d metrics", len(analytics))


# Global service instance
analytics_report_generator = AnalyticsReportGenerator()

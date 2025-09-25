"""
Test utilities for Excel processing functionality.
"""

import io
from typing import Any, Dict, List

from openpyxl.workbook import Workbook


def create_test_excel_file(data: List[Dict[str, Any]]) -> bytes:
    """
    Create a test Excel file with the provided data.

    Args:
        data: List of dictionaries representing Excel rows

    Returns:
        bytes: Excel file content as bytes
    """
    # Create workbook and worksheet
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    if not data:
        # Create empty workbook
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # Get headers from first row
    headers = list(data[0].keys())

    # Write headers
    for col, header in enumerate(headers, 1):
        worksheet.cell(row=1, column=col, value=header)

    # Write data rows
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, "")
            worksheet.cell(row=row_idx, column=col_idx, value=value)

    # Save to bytes
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_sample_excel_data() -> List[Dict[str, Any]]:
    """
    Create sample Excel data for testing.

    Returns:
        List[Dict]: Sample Excel data
    """
    return [
        {
            "SEQ": 1,
            "Category": "Name Category",
            "Name": "TestName1",
            "Rationale": "Reason1",
            "Katakana": "カタカナ1",
            "Group1": "",
            "Group2": "",
        },
        {
            "SEQ": 2,
            "Category": "Name Category",
            "Name": "TestName2",
            "Rationale": "Reason2",
            "Katakana": "",
            "Group1": "GroupA",
            "Group2": "",
        },
        {
            "SEQ": 3,
            "Category": "Name Category",
            "Name": "TestName3",
            "Rationale": "Reason3",
            "Katakana": "",
            "Group1": "GroupA",
            "Group2": "",
        },
        {
            "SEQ": "",
            "Category": "New Category",
            "Name": "",
            "Rationale": "",
            "Katakana": "",
            "Group1": "",
            "Group2": "",
        },
        {
            "SEQ": 4,
            "Category": "New Category",
            "Name": "NewName1",
            "Rationale": "NewReason1",
            "Katakana": "",
            "Group1": "",
            "Group2": "GroupB",
        },
        {
            "SEQ": 5,
            "Category": "New Category",
            "Name": "NewName2",
            "Rationale": "NewReason2",
            "Katakana": "カタカナ2",
            "Group1": "",
            "Group2": "GroupB",
        },
    ]


def validate_slides_structure(slides: List[Dict[str, Any]]) -> List[str]:
    """
    Validate the structure of generated slides.

    Args:
        slides: List of slide dictionaries to validate

    Returns:
        List[str]: List of validation error messages (empty if valid)
    """
    errors = []

    if not slides:
        errors.append("No slides generated")
        return errors

    required_fields = ["$id", "SlideType", "SlideNumber"]
    valid_slide_types = [
        "NameEvaluation",
        "GroupNameEvaluation",
        "Preview",
        "NameSummary",
    ]

    for i, slide in enumerate(slides):
        # Check required fields
        for field in required_fields:
            if field not in slide:
                errors.append(f"Slide {i + 1}: Missing required field '{field}'")

        # Check slide type
        if "SlideType" in slide and slide["SlideType"] not in valid_slide_types:
            errors.append(f"Slide {i + 1}: Invalid SlideType '{slide['SlideType']}'")

        # Check ID uniqueness
        slide_id = slide.get("$id", "")
        for j, other_slide in enumerate(slides):
            if i != j and other_slide.get("$id", "") == slide_id:
                errors.append(f"Slides {i + 1} and {j + 1}: Duplicate ID '{slide_id}'")

    # Check if last slide is NameSummary
    if slides and slides[-1].get("SlideType") != "NameSummary":
        errors.append("Last slide should be of type 'NameSummary'")

    return errors


# Sample usage examples
SAMPLE_EXCEL_FILES = {
    "basic_example": [
        {
            "SEQ": 1,
            "Category": "Name Category",
            "Name": "TestName1",
            "Rationale": "Reason1",
        },
        {
            "SEQ": 2,
            "Category": "Name Category",
            "Name": "TestName2",
            "Rationale": "Reason2",
            "Group1": "GroupA",
        },
    ],
    "katakana_example": [
        {
            "SEQ": 1,
            "Category": "Name Category",
            "Name": "TestName1",
            "Rationale": "Reason1",
            "Katakana": "カタカナ",
        },
    ],
    "grouping_example": [
        {
            "SEQ": 1,
            "Category": "Name Category",
            "Name": "Name1",
            "Rationale": "Reason1",
            "Group1": "GroupA",
        },
        {
            "SEQ": 2,
            "Category": "Name Category",
            "Name": "Name2",
            "Rationale": "Reason2",
            "Group1": "GroupA",
        },
        {
            "SEQ": 3,
            "Category": "Name Category",
            "Name": "Name3",
            "Rationale": "Reason3",
            "Group2": "GroupB",
        },
    ],
    "category_change_example": [
        {
            "SEQ": 1,
            "Category": "First Category",
            "Name": "Name1",
            "Rationale": "Reason1",
        },
        {"Category": "Second Category"},
        {
            "SEQ": 2,
            "Category": "Second Category",
            "Name": "Name2",
            "Rationale": "Reason2",
        },
    ],
}
